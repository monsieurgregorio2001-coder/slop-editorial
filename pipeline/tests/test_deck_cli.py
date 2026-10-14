import io
import re
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from pptx import Presentation

from slopgen import cli
from slopgen.deck import build_pdf, build_pptx, fit_body, image_fit, parse_slides
from slopgen.errors import MissingSecret
from slopgen.pack import load_plan
from slopgen.secrets import get_secret
from .helpers import PLAN, png


class DeckTests(unittest.TestCase):
    def setUp(self):
        self.plan = load_plan(PLAN)
        self.slides = parse_slides(self.plan.deck["markdown"], self.plan.deck["heading"])

    def test_five_slides_text_verbatim(self):
        self.assertEqual([s.number for s in self.slides], [1, 2, 3, 4, 5])
        md = self.plan.deck["markdown"]
        for s in self.slides:
            self.assertIn(s.title, md)
            for line in s.body:
                self.assertIn(line, md)
        self.assertEqual(self.slides[0].title, "Slop Editorial")
        self.assertEqual(self.slides[1].title, "The Endless Feed")

    def test_pptx_has_all_text_and_only_given_images(self):
        with tempfile.TemporaryDirectory() as d:
            out = build_pptx(self.slides, {2: png(size=(1024, 1280)), 3: png(size=(1024, 1280))}, {2, 3, 4}, Path(d) / "x.pptx")
            prs = Presentation(out)
            self.assertEqual(len(prs.slides), 5)
            pics = [sum(1 for sh in sl.shapes if sh.shape_type == 13) for sl in prs.slides]
            self.assertEqual(pics, [0, 1, 1, 0, 0])
            alltext = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
            for s in self.slides:
                for line in s.body:
                    self.assertIn(line, alltext)
            self.assertIn("Image placeholder 4:5", [sh.name for sh in prs.slides[3].shapes])

    def test_pdf_has_five_pages(self):
        with tempfile.TemporaryDirectory() as d:
            out = build_pdf(self.slides, {2: png(size=(1024, 1280))}, {2, 3, 4}, Path(d) / "x.pdf")
            data = out.read_bytes()
            self.assertTrue(data.startswith(b"%PDF"))
            self.assertEqual(len(re.findall(rb"/Type /Page\b", data)), 5)

    def test_image_is_contained_not_cropped(self):
        w, h = image_fit(png(size=(1000, 1000)), 4.0, 5.0)
        self.assertAlmostEqual(w, 4.0)
        self.assertLessEqual(h, 5.0)

    def test_fit_body_shrinks_for_long_text(self):
        self.assertGreater(fit_body(["hello"], 6.5, 4.0), fit_body(["word " * 400], 6.5, 4.0))


class CliTests(unittest.TestCase):
    def run_cli(self, *args):
        buf = io.StringIO()
        with redirect_stdout(buf):
            rc = cli.main(list(args))
        return rc, buf.getvalue()

    def test_end_to_end_with_mock_provider(self):
        with tempfile.TemporaryDirectory() as d:
            rc, out = self.run_cli("all", "--runs-dir", d, "--only", "P1,P2,P3,P4")
            self.assertEqual(rc, 0, out)
            base = Path(d) / load_plan(PLAN).name
            self.assertTrue((base / "review.html").is_file())
            self.assertTrue((base / "deck" / "review-deck.pptx").is_file())
            self.assertTrue((base / "deck" / "review-deck.pdf").is_file())
            rc, out = self.run_cli("select", "--runs-dir", d, "P1", "2")
            self.assertIn("P2", out)

    def test_plan_spends_nothing_and_lists_cost(self):
        rc, out = self.run_cli("plan", "--provider", "fal", "--runs-dir", tempfile.mkdtemp())
        self.assertEqual(rc, 0)
        self.assertIn("estimated new spend", out)

    def test_paid_run_refused_without_confirmation(self):
        rc = cli.main(["run", "--provider", "fal", "--runs-dir", tempfile.mkdtemp(), "--only", "P1"])
        self.assertEqual(rc, 1)

    def test_no_relaxed_filter_or_video_flags_exist(self):
        with self.assertRaises(SystemExit):
            cli.main(["run", "--relaxed-filter"])


class SecretTests(unittest.TestCase):
    def test_env_then_keychain_then_error(self):
        self.assertEqual(get_secret("K", environ={"K": " v "}), "v")

        class R:
            returncode, stdout = 0, "fromkeychain\n"

        self.assertEqual(get_secret("K", environ={}, runner=lambda *a, **k: R(), platform="darwin"), "fromkeychain")
        with self.assertRaises(MissingSecret) as cm:
            get_secret("K", environ={}, platform="linux")
        self.assertIn("security add-generic-password", str(cm.exception))


if __name__ == "__main__":
    unittest.main()
